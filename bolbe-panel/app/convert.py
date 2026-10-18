"""Yüklenen dosyaları ajanların okuyabileceği Markdown metnine çevirir.

Desteklenenler: metin dosyaları, PDF, Word (.docx), Excel (.xlsx/.xlsm), PowerPoint (.pptx).
Sonuç dosyanın değişme zamanına göre önbelleğe alınır.
"""

from __future__ import annotations

import re
import zipfile
from functools import lru_cache
from pathlib import Path
from xml.etree import ElementTree as ET

from fastapi import HTTPException

from . import storage

CONVERTIBLE = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx"}
MAX_ROWS, MAX_COLS = 500, 30
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def readable(path: Path) -> bool:
    return path.suffix.lower() in storage.TEXT_EXT | CONVERTIBLE


def _cell(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ").strip()


def _table(rows: list[list[str]]) -> str:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return "_(boş)_"
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        pages.append(f"## Sayfa {i}\n\n{text or '_(metin yok; taranmış görsel olabilir)_'}")
    return "\n\n".join(pages)


def _docx(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(f"{W}body")
    out: list[str] = []
    for el in list(body) if body is not None else []:
        if el.tag == f"{W}p":
            text = "".join(t.text or "" for t in el.iter(f"{W}t")).strip()
            if not text:
                continue
            style = el.find(f"{W}pPr/{W}pStyle")
            level = re.search(r"(\d)", style.get(f"{W}val", "")) if style is not None and "eading" in style.get(f"{W}val", "") else None
            bullet = el.find(f"{W}pPr/{W}numPr") is not None
            out.append(("#" * int(level.group(1)) + " " if level else "- " if bullet else "") + text)
        elif el.tag == f"{W}tbl":
            rows = [[_cell("".join(t.text or "" for t in tc.iter(f"{W}t"))) for tc in tr.iter(f"{W}tc")]
                    for tr in el.iter(f"{W}tr")]
            out.append(_table(rows))
    return "\n\n".join(out)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                rows.append([f"… ({MAX_ROWS} satırdan sonrası kesildi)"])
                break
            rows.append([_cell(v) for v in row[:MAX_COLS]])
        parts.append(f"## Sayfa: {ws.title}\n\n{_table(rows)}")
    wb.close()
    return "\n\n".join(parts)


def _pptx(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        slides = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                        key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
        parts = []
        for i, name in enumerate(slides, 1):
            root = ET.fromstring(z.read(name))
            paras = ["".join(t.text or "" for t in p.iter(f"{A}t")).strip() for p in root.iter(f"{A}p")]
            parts.append(f"## Slayt {i}\n\n" + "\n".join(f"- {p}" for p in paras if p))
    return "\n\n".join(parts)


@lru_cache(maxsize=64)
def _convert(path_str: str, mtime: float, size: int) -> str:
    path = Path(path_str)
    suffix = path.suffix.lower()
    if suffix in storage.TEXT_EXT:
        return path.read_text("utf-8", errors="replace")
    try:
        body = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx, ".pptx": _pptx}[suffix](path)
    except (zipfile.BadZipFile, KeyError, ET.ParseError, ValueError, OSError) as exc:
        raise HTTPException(422, f"Dosya dönüştürülemedi ({path.name}): {exc}") from exc
    return f"# {path.name}\n\n{body}"


def to_markdown(area: str, name: str, limit: int = 80_000) -> str:
    path = storage.area_file(area, name)
    suffix = path.suffix.lower()
    if suffix in {".xls", ".doc", ".ppt"}:
        raise HTTPException(400, f"Eski {suffix} biçimi okunamıyor; dosyayı {suffix}x olarak kaydedip yeniden yükleyin")
    if not readable(path):
        raise HTTPException(400, f"Bu dosya türü metne çevrilemiyor ({suffix or 'uzantısız'}); görsel ise view_image kullanın")
    st = path.stat()
    text = _convert(str(path), st.st_mtime, st.st_size)
    return text if len(text) <= limit else text[:limit] + f"\n\n… (metin {limit} karakterde kesildi)"
